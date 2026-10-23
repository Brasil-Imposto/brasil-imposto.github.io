"""Judiciary 2026: ministers of STF, STJ, TST, STM and TSE, pay by component from official files.

Raw files live in data/federal/raw-pay/judiciary/ (see NOTES.md there for how each was fetched):
  stf/stf-2026-<m>.html    STF eGesp "Rendimentos da folha", Situação Funcional = MINISTRO (court's own file)
  cnj/<court>-2026-all-<tab>.xlsx  CNJ panel Anexo VIII (Portaria CNJ 63/2017), each court's uploaded spreadsheet
  tst/tst-2026-<mm>.csv    TST own remuneration CSV (used for names, roles and a cross-check)
  comp/*.html              composition pages (STF, STJ, TSE via Web Archive)
If cnj/*.xlsx are missing, run the CNJ panel export first:
  uv run python -B docs/superpowers/tools/collect-judiciary-cnj-download.py
Run: uv run python -B docs/superpowers/tools/collect-judiciary-2026.py
"""
import collections, glob, html, json, re, subprocess, unicodedata
from urllib.parse import quote
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
RAW = ROOT / "data/federal/raw-pay/judiciary"
OUT = ROOT / "data/federal/judiciary-2026.json"
TODAY = "2026-10-21"
MES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]

STF_URL = "https://egesp-portal.stf.jus.br/transparencia/rendimento_folha"
STF_COMP = "https://portal.stf.jus.br/textos/verTexto.asp?servico=bibliotecaConsultaProdutoBibliotecaPastaMinistro&pagina=ComposicaoAtual"
CNJ_URL = "https://paineis.cnj.jus.br/QvAJAXZfc/opendoc.htm?document=qvw_l%2FPainelCNJ.qvw&host=QVS%40neodimio03&anonymous=true&sheet=shPORT63Relatorios"
TST_URL = "https://transparencia-remuneracao.tst.jus.br/"
STJ_LIST = "https://www.stj.jus.br/web/verMinistrosSTJ?parametro=1"
TSE_LIVE = "https://www.tse.jus.br/institucional/ministros"
TSE_ARCH = "https://web.archive.org/web/20260903124144/https://www.tse.jus.br/institucional/ministros"
STM_ATA44 = "https://dspace.stm.jus.br/handle/123456789/206325"
STM_ATA37 = "https://dspace.stm.jus.br/handle/123456789/205914"
STJ_ORIGIN = "https://www.stj.jus.br/sites/portalp/SiteAssets/Paginas/Institucional/Composicao/Classe_de_origem_ativos.pdf"
STJ_COMP_PDF = "https://www.stj.jus.br/sites/portalp/SiteAssets/Paginas/Institucional/Composicao/Composicao-do-STJ.pdf"
TST_MIN = "https://www.tst.jus.br/ministros"
STF_DECREE_NEWS = "https://noticias.stf.jus.br/postsnoticias/publicado-decreto-de-aposentadoria-do-ministro-luis-roberto-barroso/"
CNJ_PC3_2025 = "https://atos.cnj.jus.br/atos/detalhar/6023"
PC1_2026_TRIED = ["https://juslaboris.tst.jus.br/handle/20.500.12178/260055",
                  "https://juslaboris.tst.jus.br/bitstream/handle/20.500.12178/260055/2026_portc0001_cnj_etc.pdf?sequence=1&isAllowed=y",
                  "https://www.tse.jus.br/legislacao/compilada/prtc/2026/portaria-conjunta-no-1-de-3-de-fevereiro-de-2026"]
PC3_VALUE = 1784.42
LAW_14520 = "https://www.planalto.gov.br/ccivil_03/_ato2023-2026/2023/lei/L14520.htm"
CF = "https://www.planalto.gov.br/ccivil_03/constituicao/constituicao.htm"
STF_SALARY, HIGHER_SALARY = 46366.19, 44047.88

STOP = {"DE", "DA", "DO", "DOS", "DAS", "E"}


def norm_tokens(s):
    s = unicodedata.normalize("NFKD", s.replace("​", "")).encode("ascii", "ignore").decode().upper()
    return [t for t in re.findall(r"[A-Z]+", s) if t not in STOP]


def slug(s):
    return "-".join(norm_tokens(s)).lower()


def match(short, candidates):
    st = norm_tokens(short)
    hits = [c for c in candidates if all(t in norm_tokens(c) for t in st)]
    if len(hits) > 1:
        hits.sort(key=lambda c: len(norm_tokens(c)))
    return hits


def num(v):
    if v in (None, "", "0"):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    v = str(v).replace("R$", "").replace(".", "").replace(",", ".").strip()
    try:
        return float(v)
    except ValueError:
        return 0.0


def months_label(ms):
    ms = sorted(ms)
    if not ms:
        return None
    if ms == list(range(ms[0], ms[-1] + 1)) and len(ms) > 2:
        return f"{MES[ms[0]-1]} a {MES[ms[-1]-1]} 2026"
    return ", ".join(MES[m - 1] for m in ms) + " 2026"


def text_of(t):
    t = re.sub(r"<script.*?</script>|<style.*?</style>", "", t, flags=re.S)
    t = html.unescape(re.sub(r"<[^>]+>", "\n", t))
    return [x.strip() for x in t.split("\n") if x.strip()]


# ---------- compositions ----------
def stf_composition():
    t = (RAW / "comp/stf-comp.html").read_text(errors="ignore")
    photos = [p for p in re.findall(r'src="(http://www\.stf\.jus\.br/arquivo/cms/bibliotecaConsultaProdutoBibliotecaPastaMinistro/imagem/foto_[^"]+)"', t)]
    lines = text_of(t)
    people = []
    for i, l in enumerate(lines):
        if l in ("Ministro", "Ministra") and i + 1 < len(lines):
            name = lines[i + 1]
            role = lines[i + 2].strip("()") if i + 2 < len(lines) and lines[i + 2].startswith("(") else l
            people.append({"short": name, "role": role})
        if l.startswith("Essa informa"):
            break
    if not (len(people) == len(photos) == 10):
        raise SystemExit((len(people), len(photos)))
    for p, ph in zip(people, photos):
        p["photo"] = ph.strip().replace("http://", "https://")
    return people


def stj_composition():
    t = (RAW / "comp/stj-ver.html").read_text(errors="ignore")
    out = []
    for href, inner in re.findall(r'<span class="clsMinistrosNome"><a href="([^"]+)">(.*?)</a></span>', t, re.S):
        txt = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", "", inner))).strip()
        name, _, role = txt.partition(" - ")
        out.append({"short": name.strip(), "role": role.strip() or "Ministro",
                    "page": "https://www.stj.jus.br/web/" + html.unescape(href)})
    return out


def tse_composition():
    lines = text_of((RAW / "comp/tse-arch.html").read_text(errors="ignore"))
    i = lines.index("MINISTROS EFETIVOS")
    j = lines.index("MINISTROS SUBSTITUTOS")
    end = next(k for k in range(j, len(lines)) if lines[k].startswith("A Corte"))
    out = []
    for block, kind in ((lines[i + 5:j], "efetivo"), (lines[j + 5:end], "substituto")):
        k = 0
        while k < len(block):
            name = block[k]
            if name == "Vago":
                k += 2
                continue
            role = None
            if block[k + 1].startswith("("):
                role = block[k + 1].strip("()")
                k += 1
            origin, start, finish, bien = block[k + 1:k + 5]
            out.append({"short": name, "tse_office": role, "kind": kind, "origin": origin,
                        "start": start, "end": finish, "bienio": bien})
            k += 5
    return out


def pdf_text(path):
    return subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True, check=True).stdout


def stm_composition():
    t44 = re.sub(r"\s+", " ", pdf_text(RAW / "stm/ata-206325.pdf"))
    t37 = re.sub(r"\s+", " ", pdf_text(RAW / "stm/ata37.pdf"))
    pres = re.search(r"PRESIDÊNCIA DA MINISTRA Dra\. ([A-ZÀ-Ú ]+?) Presentes", t44).group(1).strip()
    present = re.search(r"Presentes (.*?)\. Presente o", t44).group(1)
    names = [n.strip() for n in re.split(r",? (?:e )?(?:o Ministro|a Ministra) ", " " + present) if n.strip()]
    vice = re.search(r"Ministro ([A-ZÀ-Ú ]+?), Vice-Presidente", t37).group(1).strip()
    corr = "Corregedor- Geral da Justiça Militar da União" in t37
    out = [{"short": " ".join(w.capitalize() if w.lower() not in STOP_LOWER else w.lower() for w in pres.split()), "role": "Presidente"}]
    for n in names:
        role = "Ministro"
        if norm_tokens(n) == norm_tokens(vice):
            role = "Vice-Presidente e Corregedor-Geral da Justiça Militar da União" if corr else "Vice-Presidente"
        out.append({"short": n, "role": role})
    if not (len(out) == 15):
        raise SystemExit(len(out))
    return out


STOP_LOWER = {"de", "da", "do", "dos", "das", "e"}


def stj_origin():
    rows = {}
    for line in pdf_text(RAW / "vac/stj-origem.pdf").splitlines():
        m = re.match(r"\s*(\d+)\s+(.+?)(\*?)\s{2,}(Min\..+?)\s{2,}(.+?)\s{2,}(\d{1,2})/(\d{1,2})/(\d{4})\s+(\d{1,2})/(\d{1,2})/(\d{4})", line)
        if m:
            g = m.groups()
            rows[g[1].strip()] = {"seat_of": g[3].strip(), "origin": re.sub("[\u2013\u2014]", "-", g[4].strip()), "available": bool(g[2]),
                                  "nominated": f"{g[7]}-{int(g[6]):02d}-{int(g[5]):02d}", "since": f"{g[10]}-{int(g[9]):02d}-{int(g[8]):02d}"}
    return rows


def stj_photo(page):
    cod = page.split("cod_matriculamin=")[-1]
    f = RAW / f"comp/stjcv/{cod}.html"
    if not f.exists():
        return None
    ph = [s for s in re.findall(r'<img[^>]*src="([^"]+)"', f.read_text(errors="ignore")) if "/ministros/fotos/" in s.lower()]
    return "https://www.stj.jus.br" + quote(html.unescape(ph[0])) if ph else None


def tst_cards():
    t = (RAW / "comp/tst-min2.html").read_text(errors="ignore")
    out = []
    for c in t.split('<h3 class="asset-title">')[1:]:
        a = re.search(r'<a href="([^"]+)"[^>]*>\s*(.*?)\s*</a>', c, re.S)
        img = re.search(r'<img[^>]*src="([^"]+)"', c)
        title = re.sub(r"^\d+ - ", "", html.unescape(re.sub(r"\s+", " ", a.group(2))).strip())
        name = title.split(" - ")[0]
        src = html.unescape(img.group(1)) if img else None
        if src and src.startswith("/"):
            src = "https://www.tst.jus.br" + src
        out.append({"name": name, "page": html.unescape(a.group(1)).split("?")[0], "photo": src})
    return out


def overlap(a, b):
    ta, tb = set(norm_tokens(a)), set(norm_tokens(b))
    return len(ta & tb) / max(1, min(len(ta), len(tb)))


# ---------- pay files ----------
STF_COLS = [
    ("Vencimentos / Subsídios", "salary", "Vencimentos / Subsídios"),
    ("Vantagens Pessoais", "personal_advantages", "Vantagens Pessoais"),
    ("Vantagens de natureza periódica/eventual ou relativas às lotações dos servidores", "eventual_advantages", "Vantagens de natureza periódica/eventual"),
    ("Exercício de cargo em comissão/função comissionada", "commission", "Exercício de cargo em comissão/função comissionada"),
    ("Abono de Permanência", "permanence_bonus", "Abono de Permanência"),
    ("Férias", "vacation", "Férias"),
    ("Gratificação natalina e antecipação", "christmas_bonus", "Gratificação natalina e antecipação"),
    ("Auxílios e benefícios", "benefits", "Auxílios e benefícios"),
    ("Indenizações", "indemnities", "Indenizações"),
    ("Exercícios Anteriores e licença prêmio convertida em pecúnia", "back_pay", "Exercícios Anteriores e licença prêmio convertida em pecúnia"),
    ("Auxílio Moradia", "housing", "Auxílio Moradia"),
]


def stf_pay():
    data = collections.defaultdict(lambda: collections.defaultdict(dict))  # name -> month -> key -> value
    months = []
    for f in sorted(glob.glob(str(RAW / "stf/stf-2026-*.html"))):
        m = int(re.search(r"stf-2026-(\d+)\.html", f).group(1))
        t = Path(f).read_text(errors="ignore")
        rows = [[html.unescape(re.sub(r"<[^>]+>", "", c)).strip() for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", r, re.S)]
                for r in re.findall(r"<tr[^>]*>(.*?)</tr>", t, re.S)]
        hdr = next(r for r in rows if r and r[0] == "Matricula")
        ci = {h: i for i, h in enumerate(hdr)}
        n = 0
        for r in rows:
            if len(r) == len(hdr) and r[0] != "Matricula" and r[ci["Situação Funcional"]] == "MINISTRO":
                n += 1
                for col, key, _ in STF_COLS:
                    data[r[ci["Nome"]]][m][key] = data[r[ci["Nome"]]][m].get(key, 0) + num(r[ci[col]])
        if n:
            months.append(m)
    return data, months


TABS = {
    "Contracheque": [("Subsídio (R$)", "salary", "Subsídio"), ("Diárias (12) (R$)", "per_diem", "Diárias")],
    "Direitos Pessoais": [("Abono de permanência (R$)", "permanence_bonus", "Abono de permanência"), ("Outra (R$)", "personal_other", "Direitos pessoais: outra")],
    "Indeniza": [("Auxílio-alimentação (R$)", "food", "Auxílio-alimentação"), ("Auxílio Pré-escolar (R$)", "preschool", "Auxílio Pré-escolar"),
                 ("Auxílio Saúde (R$)", "health", "Auxílio Saúde"), ("Auxílio Natalidade (R$)", "birth", "Auxílio Natalidade"),
                 ("Auxílio Moradia (R$)", "housing", "Auxílio Moradia"), ("Ajuda de Custo (R$)", "relocation", "Ajuda de Custo"),
                 ("Outra (R$)", "indemnity_other", "Indenizações: outra")],
    "Direitos Eventuais": [("Abono constitucional de 1/3 de férias (R$)", "vacation_third", "Abono constitucional de 1/3 de férias"),
                           ("Indenização de férias (R$)", "vacation_indemnity", "Indenização de férias"),
                           ("Antecipação de férias (R$)", "vacation_advance", "Antecipação de férias"),
                           ("Gratificação natalina (R$)", "christmas_bonus", "Gratificação natalina"),
                           ("Antecipação de gratificação natalina (R$)", "christmas_advance", "Antecipação de gratificação natalina"),
                           ("Substituição (R$)", "substitution", "Substituição"),
                           ("Gratificação por exercício cumulativo (R$)", "cumulative_jurisdiction", "Gratificação por exercício cumulativo"),
                           ("Gratificação por encargo Curso/Concurso (R$)", "course_bonus", "Gratificação por encargo Curso/Concurso"),
                           ("Pagamentos retroativos (R$)", "back_pay", "Pagamentos retroativos"),
                           ("JETON (R$)", "jeton", "JETON"), ("Outra (R$)", "eventual_other", "Direitos eventuais: outra")],
}
CNJ_ORIGIN = ("Remuneração do órgão de origem (11) (R$)", "origin_pay", "Remuneração do órgão de origem")
CNJ_GROSS = ("Total de Rendimentos (4)", "_gross", "Total de Rendimentos")


def cnj_pay(court):
    data = collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(float)))
    details = collections.defaultdict(lambda: collections.defaultdict(set))
    months = set()
    for tab, cols in TABS.items():
        cols = cols + ([CNJ_ORIGIN, CNJ_GROSS] if tab == "Contracheque" else [])
        ws = openpyxl.load_workbook(RAW / f"cnj/{court}-2026-all-{tab}.xlsx", read_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        hdr = rows[0]
        for r in rows[1:]:
            if r[0] != court.upper() or not r[2]:
                continue
            m = int(str(r[2])[:2])
            months.add(m)
            for i, h in enumerate(hdr):
                for col, key, _ in cols:
                    if h == col:
                        data[r[1]][m][key] += num(r[i])
                        if col == "Outra (R$)" and num(r[i]) and i + 1 < len(hdr) and hdr[i + 1] == "Detalhe" and r[i + 1] not in (None, "0"):
                            details[r[1]][key].add(str(r[i + 1]).strip())
    return data, sorted(months), details


def control_files(court):
    ws = openpyxl.load_workbook(RAW / f"cnj/{court}-2026-all-Controle de Arquivos.xlsx", read_only=True).worksheets[0]
    return [r[2] for r in list(ws.iter_rows(values_only=True))[1:] if r[2]]


def tst_csv():
    out = {}
    for f in sorted(glob.glob(str(RAW / "tst/tst-2026-*.csv"))):
        m = int(re.search(r"tst-2026-(\d+)\.csv", f).group(1))
        for line in Path(f).read_bytes().decode("latin-1").splitlines():
            c = line.split(";")
            if len(c) > 8 and c[1] == "MINISTRO" and c[2] != "INATIVO":
                out.setdefault(c[0], {"lot": c[2], "gross": {}, "last": m})["gross"][m] = num(c[8])
                out[c[0]]["lot"], out[c[0]]["last"] = c[2], m
    latest = max(v["last"] for v in out.values())
    return {k: v for k, v in out.items() if v["last"] == latest}


# ---------- assembly ----------
def build_items(pay, months, cols, source, details=None, prefix="", label_prefix="", note_extra=""):
    items = []
    for key, label in cols:
        vals = {m: pay.get(m, {}).get(key, 0.0) for m in months}
        present = [m for m in months if m in pay]
        total = round(sum(vals.values()), 2)
        note = f"Soma dos meses publicados; média = total / {len(months)} meses publicados pelo tribunal."
        if details and details.get(key):
            note += " Detalhe informado: " + "; ".join(sorted(details[key]))[:300] + "."
        if note_extra:
            note += " " + note_extra
        if len(present) < len(months):
            note += f" A pessoa aparece em {len(present)} dos {len(months)} meses."
        items.append({"key": prefix + key, "label": label_prefix + label, "uses": total > 0,
                      "monthly": round(total / len(months), 2) if months else None,
                      "year_total": total, "months": months_label(months), "note": note, "source": source})
    return items


def as_of(pay):
    return f"2026-{max(pay):02d}" if pay else None


def main():
    sources, people, positions, not_found = [], [], [], []
    src = lambda n, u: sources.append({"name": n, "url": u, "read_on": TODAY})
    src("STF eGesp, Rendimentos da folha (Situação Funcional: Ministro), jan a ago 2026", STF_URL)
    src("STF, Composição atual", STF_COMP)
    src("CNJ, Painel Portaria 63/2017 (Anexo VIII, Res. CNJ 102/2009 e 215/2015): STJ, TST, STM, TSE", CNJ_URL)
    src("TST, Remuneração (Lei de Acesso à Informação), CSV mensal", TST_URL)
    src("STJ, Ministros em atividade", STJ_LIST)
    src("TSE, Composição da Corte (cópia do Web Archive de 2026-09-03; o site do TSE bloqueia acesso automatizado)", TSE_ARCH)
    src("Lei 14.520/2023 (subsídio dos ministros do STF)", LAW_14520)
    src("Constituição Federal, art. 93, V", CF)

    # STF
    comp = stf_composition()
    pay, stf_months = stf_pay()
    stf_names = list(pay)
    stf_by_short = {}
    for p in comp:
        hit = match(p["short"], stf_names)
        full = hit[0] if hit else None
        stf_by_short[p["short"]] = full
        adm = None
        for f in glob.glob(str(RAW / "stf/det-[0-9]*.html")):
            t = re.sub(r"(\s*\|\s*)+", "|", html.unescape(re.sub(r"<[^>]+>", "|", Path(f).read_text(errors="ignore").replace("\\n", " "))))
            g = re.search(r"\|Nome\|([^|]+)\|Data de Admiss[^|]*\|(\d\d)/(\d\d)/(\d{4})", t)
            if g and full and g.group(1).strip() == full:
                adm = f"{g.group(4)}-{g.group(3)}-{g.group(2)}"
        people.append({"id": f"stf:{slug(p['short'])}", "court": "STF", "position": "ministros-stf", "name": p["short"],
                       "full_name": full, "role": p["role"], "tse_role": None, "since": adm, "photo": p["photo"],
                       "photo_credit": "STF", "page": STF_COMP,
                       "items": build_items(pay.get(full, {}), stf_months, [(k, l) for _, k, l in STF_COLS], STF_URL,
                                            note_extra="O STF publica Auxílios e benefícios e Indenizações em uma linha cada, sem detalhar."),
                       "data_as_of": as_of(pay.get(full, {}))})
    positions.append({"id": "ministros-stf", "court": "STF", "count": len(comp), "items": [
        {"key": "salary", "label": "Subsídio", "value_monthly": STF_SALARY,
         "note": "Desde 2025-02-01. A página de composição lista 10 ministros; o tribunal tem 11 cadeiras.", "source": LAW_14520}]})

    # CNJ courts
    cnj_cols = [(k, l) for tab in TABS.values() for _, k, l in tab] + [(CNJ_ORIGIN[1], CNJ_ORIGIN[2])]
    cnj = {c: cnj_pay(c) for c in ("stj", "tst", "stm", "tse")}
    tse_comp = tse_composition()
    tse_data, tse_months, tse_det = cnj["tse"]

    def tse_for(short, full):
        hits = [t for t in tse_comp if match(t["short"], [full or short]) or match(short, [t["short"]])]
        return hits[0] if hits else None

    def tse_items(full_names):
        names = [n for n in tse_data if any(match(fn, [n]) or match(n, [fn]) for fn in full_names if fn)]
        if not names:
            return [], None
        pay = tse_data[names[0]]
        keep = [(k, l) for k, l in cnj_cols if k != "origin_pay" and round(sum(pay.get(m, {}).get(k, 0) for m in tse_months), 2) > 0]
        return build_items(pay, tse_months, keep, CNJ_URL, tse_det.get(names[0]), "tse_", "TSE: ",
                           "Pagamento do TSE (arquivo TSE no painel do CNJ)."), names[0]

    def role_text(t):
        office = f"{t['tse_office']} do TSE" if t["tse_office"] else f"Ministro {t['kind']} do TSE"
        txt = f"{office} ({t['kind']}, origem {t['origin']}, de {t['start']} a {t['end']}, {t['bienio']} biênio)"
        if t["end"] == "19.9.2026":
            txt += "; o biênio terminou em 19.9.2026 e a composição posterior não foi confirmada"
        return txt

    for p in people:
        t = tse_for(p["name"], p["full_name"])
        if t:
            p["tse_role"] = role_text(t)
            extra, _ = tse_items([p["full_name"], p["name"]])
            p["items"] += extra

    # STJ
    stj_data, stj_months, stj_det = cnj["stj"]
    stj_comp = stj_composition()
    origin = stj_origin()
    if not (len(origin) == len(stj_comp)):
        raise SystemExit((len(origin), len(stj_comp)))
    stj_found = 0
    for p in stj_comp:
        hit = match(p["short"], list(stj_data))
        full = hit[0] if hit else None
        stj_found += bool(full)
        pr = {"id": f"stj:{slug(p['short'])}", "court": "STJ", "position": "ministros-stj", "name": p["short"], "full_name": full,
              "role": p["role"], "tse_role": None, "since": None, "photo": None, "photo_credit": None, "page": p["page"],
              "items": build_items(stj_data.get(full, {}), stj_months, cnj_cols, CNJ_URL, stj_det.get(full)) if full else [],
              "data_as_of": as_of(stj_data.get(full, {})) if full else None}
        if not full:
            pr["note"] = "Nome não encontrado no arquivo do STJ no painel do CNJ."
        orow = max(origin.items(), key=lambda kv: overlap(kv[0], p["short"]))
        if overlap(orow[0], p["short"]) >= 0.75:
            o = orow[1]
            pr["since"] = o["since"]
            pr["since_note"] = f"Posse no STJ; nomeação em {o['nominated']}; vaga do {o['seat_of']}; origem: {o['origin']} (tabela de antiguidade do STJ, 19/8/2026)."
            if o["available"]:
                pr["note"] = (pr.get("note", "") + " Em disponibilidade, segundo a tabela de antiguidade do STJ de 19/8/2026.").strip()
        pr["photo"] = stj_photo(p["page"])
        pr["photo_credit"] = "STJ" if pr["photo"] else None
        t = tse_for(p["short"], full)
        if t:
            pr["tse_role"] = role_text(t)
            extra, _ = tse_items([full, p["short"]])
            pr["items"] += extra
        people.append(pr)
    positions.append({"id": "ministros-stj", "court": "STJ", "count": len(stj_comp), "items": [
        {"key": "salary", "label": "Subsídio", "value_monthly": HIGHER_SALARY, "note": f"95% do subsídio do STF (CF art. 93, V; Lei 14.520/2023). A lista oficial traz {len(stj_comp)} ministros em atividade; o tribunal tem 33 cadeiras.", "source": CF}]})

    # TST
    tst_data, tst_months, tst_det = cnj["tst"]
    own = tst_csv()
    cards = tst_cards()
    mism = checked = 0
    mism_months = collections.Counter()
    for full, info in own.items():
        lot = info["lot"]
        role = ("Presidente" if "Presidência" in lot and "Vice" not in lot else "Vice-Presidente" if "Vice-Presid" in lot
                else "Corregedor-Geral da Justiça do Trabalho" if "Corregedoria" in lot else "Ministro")
        hit = [n for n in tst_data if norm_tokens(n) == norm_tokens(full)] or match(full, list(tst_data))
        cn = hit[0] if hit else None
        for m, v in info["gross"].items():
            checked += 1
            if not cn or abs(tst_data[cn].get(m, {}).get("_gross", 0) - v) > 0.01:
                mism += 1
                mism_months[m] += 1
        name = " ".join(w.capitalize() if w.lower() not in ("de", "da", "do", "dos", "das", "e") else w.lower() for w in full.split())
        pr = {"id": f"tst:{slug(full)}", "court": "TST", "position": "ministros-tst", "name": name, "full_name": full, "role": role,
              "tse_role": None, "since": None, "photo": None, "photo_credit": None, "page": TST_URL,
              "items": build_items(tst_data.get(cn, {}), tst_months, cnj_cols, CNJ_URL, tst_det.get(cn)) if cn else [],
              "data_as_of": as_of(tst_data.get(cn, {})) if cn else None}
        if "DISPOSI" in lot.upper():
            pr["note"] = f"Lotação no arquivo do TST: {lot}."
        card = max(cards, key=lambda c: overlap(c["name"], full))
        if overlap(card["name"], full) >= 0.75:
            pr["photo"], pr["photo_credit"], pr["page"] = card["photo"], "TST", card["page"]
        else:
            pr["note"] = (pr.get("note", "") + " Não aparece na página Ministros do TST (2026-10-21).").strip()
        people.append(pr)
    print(f"TST cross-check: own CSV Total de Rendimentos vs CNJ Total de Rendimentos, {checked} person-months checked, {mism} differ by more than R$ 0,01")
    if mism:
        not_found.append({"what": "TST: conferência entre o CSV do próprio TST e o arquivo do TST no painel do CNJ",
                          "why": f"Total de Rendimentos difere em {mism} de {checked} pessoa-mês (por mês: " + ", ".join(f"{MES[m-1]} {n}" for m, n in sorted(mism_months.items())) + "). Nesses meses o CSV do TST soma mais que o Anexo VIII enviado ao CNJ; a causa (folha suplementar ou rubrica fora do Anexo VIII) não foi confirmada. Os itens por componente usam o Anexo VIII do painel do CNJ.",
                          "tried": [TST_URL, CNJ_URL]})
    positions.append({"id": "ministros-tst", "court": "TST", "count": len(own), "items": [
        {"key": "salary", "label": "Subsídio", "value_monthly": HIGHER_SALARY, "note": "95% do subsídio do STF (CF art. 93, V; Lei 14.520/2023).", "source": CF}]})

    # STM
    stm_data, stm_months, stm_det = cnj["stm"]
    for p in stm_composition():
        hit = match(p["short"], list(stm_data))
        full = hit[0] if hit else None
        people.append({"id": f"stm:{slug(p['short'])}", "court": "STM", "position": "ministros-stm", "name": p["short"], "full_name": full,
                       "role": p["role"], "tse_role": None, "since": None, "photo": None, "photo_credit": None, "page": STM_ATA44,
                       "items": build_items(stm_data.get(full, {}), stm_months, cnj_cols, CNJ_URL, stm_det.get(full)) if full else [],
                       "data_as_of": as_of(stm_data.get(full, {})) if full else None})
        if not full:
            people[-1]["note"] = "Nome não encontrado no arquivo do STM no painel do CNJ."
    positions.append({"id": "ministros-stm", "court": "STM", "count": 15, "items": [
        {"key": "salary", "label": "Subsídio", "value_monthly": HIGHER_SALARY, "note": "95% do subsídio do STF (CF art. 93, V; Lei 14.520/2023). Contagem: 15 cadeiras (CF art. 123); os 15 constam da ata da 44ª sessão de julgamento do STM, 29/9/2026.", "source": CF}]})

    # TSE jurists
    for t in tse_comp:
        if t["origin"] != "JURI":
            continue
        extra, nm = tse_items([t["short"]])
        items = [dict(i, key=i["key"][4:], label=i["label"][5:]) for i in extra]
        people.append({"id": f"tse:{slug(t['short'])}", "court": "TSE", "position": "ministros-tse", "name": t["short"], "full_name": nm,
                       "role": f"Ministro {t['kind']} (classe dos juristas)", "tse_role": role_text(t), "since": None,
                       "photo": None, "photo_credit": None, "page": TSE_ARCH, "items": items,
                       "data_as_of": as_of(tse_data.get(nm, {})) if nm else None})

    # position-level allowances fixed by acts: not collected
    for court, (data, months, _) in cnj.items():
        if court == "tse":
            continue
        pos = next(p for p in positions if p["court"] == court.upper())
        for key, label in (("food", "Auxílio-alimentação"), ("health", "Auxílio Saúde")):
            vals = collections.Counter(round(data[n][m].get(key, 0), 2) for n in data for m in data[n] if data[n][m].get(key, 0) > 0)
            mode = vals.most_common(1)[0][0] if vals else None
            seen = f" Valor pago mais frequente no arquivo do tribunal em 2026: R$ {mode:.2f} (observado, não é o valor do ato)." if mode else ""
            if key == "food":
                act = ("Portaria Conjunta nº 3, de 17/03/2025 (CNJ, TSE, STJ e CJF, TST e CSJT, STM, TJDFT): R$ 1.784,42 a partir de 1º/4/2025; "
                       "a implantação em cada órgão depende de declaração de disponibilidade orçamentária.")
                same = mode is not None and abs(mode - PC3_VALUE) < 0.005
                pos["items"].append({"key": key, "label": label, "value_monthly": PC3_VALUE if same else None,
                                     "note": act + seen + ("" if same else " O valor pago difere do ato lido; o ato posterior que o fixa não pôde ser aberto."),
                                     "source": CNJ_PC3_2025})
            else:
                pos["items"].append({"key": key, "label": label, "value_monthly": None,
                                     "note": "Ato que fixa o valor não localizado." + seen, "source": CNJ_URL})

    not_found.append({"what": "STM: fotos e data de posse dos ministros",
                      "why": "O site do STM (www.stm.jus.br) mostra 'Comunicado de Manutenção' em 2026-10-21 e o Web Archive estava fora do ar. A composição veio das atas de sessão publicadas na biblioteca digital do STM (dspace.stm.jus.br), que não trazem foto nem data de posse.",
                      "tried": ["https://www.stm.jus.br/o-stm-stm/ministros", "https://www.stm.jus.br/o-stm-stm/composicao", "https://web.archive.org/web/2026/https://www.stm.jus.br/o-stm-stm/ministros", STM_ATA44]})
    not_found.append({"what": "TST: data de posse dos ministros; STJ: foto de alguns ministros",
                      "why": "As páginas individuais dos ministros do TST respondem 'Human Verification' a acesso automatizado e a página Ministros do TST não traz datas. Algumas páginas de currículo do STJ não trazem foto. O arquivo do painel do CNJ não tem aba de dados cadastrais (abas: Contracheque, Direitos Pessoais, Indenizações, Direitos Eventuais, Controle de Arquivos), então não dá data de admissão.",
                      "tried": [TST_MIN, "https://www.tst.jus.br/ministros/-/asset_publisher/vKn1/content/15-guilherme-augusto-caputo-bastos", STJ_LIST, CNJ_URL]})
    not_found.append({"what": "TSE: fotos e composição ao vivo",
                      "why": "tse.jus.br responde 403 / Access Denied (também com navegador de mesa) e o Web Archive estava fora do ar em 2026-10-21; usada a cópia do Web Archive de 2026-09-03, sem fotos. O biênio do corregedor-geral Antonio Carlos Ferreira terminou em 19.9.2026 e a composição posterior não foi confirmada.",
                      "tried": [TSE_LIVE, "https://www.tse.jus.br/transparencia-e-prestacao-de-contas/gestao-de-pessoas", TSE_ARCH]})
    not_found.append({"what": "Auxílio-alimentação: ato de 2026 (Portaria Conjunta nº 1, de 3/2/2026, citada como fixando o valor de 2026)",
                      "why": "A biblioteca do TST (juslaboris.tst.jus.br) recusou conexão e o site do TSE responde Access Denied em 2026-10-21. O último ato lido é a Portaria Conjunta nº 3/2025 (R$ 1.784,42), que é o valor pago pelo STJ; TST e STM pagaram valores maiores em 2026.",
                      "tried": PC1_2026_TRIED + ["https://atos.cnj.jus.br/atos?ano=2026&numero=1&tipo=Portaria+Conjunta"]})
    not_found.append({"what": "Auxílio-saúde: ato que fixa o valor em cada tribunal em 2026; STF: atos de auxílios",
                      "why": "Ato não localizado. O STF não é signatário das portarias conjuntas de auxílio-alimentação e publica auxílios em uma linha só.",
                      "tried": ["https://transparencia.stj.jus.br/gestao-de-pessoas/beneficiarios-por-beneficio/", "https://dspace.stm.jus.br/discover?query=%22aux%C3%ADlio-sa%C3%BAde%22+2026"]})
    not_found.append({"what": "STJ: meses fev, mar, jul, ago e set de 2026",
                      "why": "O STJ enviou ao painel do CNJ só jan, abr, mai e jun (aba Controle de Arquivos: " + "; ".join(control_files("stj")) + "). O portal de transparência do STJ não publica folha mensal de magistrados fora do painel.",
                      "tried": ["https://transparencia.stj.jus.br/gestao-de-pessoas/pagamento-e-lotacao-de-pessoal/",
                                "https://transparencia.stj.jus.br/gestao-de-pessoas/membros-e-agentes-publicos-terceirizados-e-outros/", CNJ_URL]})
    not_found.append({"what": "STF: setembro de 2026",
                      "why": "O eGesp do STF devolve HTTP 500 para mês 9 em 2026-10-21 (ainda não publicado).", "tried": [STF_URL]})

    for p in people:
        if p["court"] in ("STJ", "TST") and not p["items"]:
            not_found.append({"what": f"{p['court']}: pagamento de {p['name']}", "why": "Nome não encontrado no arquivo do tribunal no painel do CNJ.", "tried": [CNJ_URL]})

    vacancies = [
        {"court": "STF", "position": "ministros-stf", "seat": "cadeira do ministro Luís Roberto Barroso", "since": "2025-10-18",
         "note": "Decreto de aposentadoria publicado em edição extra do DOU de 15/10/2025, com efeitos a partir de 18/10/2025. A composição atual do STF lista 10 ministros.",
         "source": STF_DECREE_NEWS},
        {"court": "STJ", "position": "ministros-stj", "seat": "vaga do Min. Antonio Palheiro (33ª cadeira na tabela de antiguidade)", "since": None,
         "note": "A tabela de antiguidade do STJ de 19/8/2026 lista a 33ª cadeira como vaga e o quadro de composição a identifica como vaga do Min. Antonio Palheiro. A data da aposentadoria não foi confirmada em fonte oficial.",
         "source": STJ_ORIGIN},
    ]
    src("STJ, Ministros ativos: antiguidade e classe de origem (19/8/2026)", STJ_ORIGIN)
    src("STJ, Composição dos colegiados (PDF)", STJ_COMP_PDF)
    src("STM, Ata da 44ª Sessão de Julgamento, 29/9/2026 (biblioteca digital do STM)", STM_ATA44)
    src("STM, Ata da 37ª Sessão de Julgamento, 1/9/2026 (vice-presidente e corregedor)", STM_ATA37)
    src("TST, Ministros", TST_MIN)
    src("CNJ, Portaria Conjunta nº 3, de 17/03/2025 (auxílio-alimentação do Poder Judiciário da União)", CNJ_PC3_2025)
    src("STF Notícias, decreto de aposentadoria do ministro Luís Roberto Barroso", STF_DECREE_NEWS)
    out = {"collected_on": TODAY, "sources": sources, "positions": positions, "people": people, "vacancies": vacancies, "not_found": not_found}
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1))
    by = collections.Counter(p["court"] for p in people)
    paid = collections.Counter(p["court"] for p in people if p["items"])
    print("people", dict(by), "with items", dict(paid))
    print("months", {"STF": stf_months, **{c.upper(): cnj[c][1] for c in cnj}})
    for p in people:
        if not p.get("full_name"):
            print("UNMATCHED", p["court"], p["name"])


if __name__ == "__main__":
    main()

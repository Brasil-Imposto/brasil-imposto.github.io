"""Tesouro base "Despesas e Transferências Totais da União" (valor pago), one year, every row, keeping the month.

Usage: .venv/bin/python agg_dtt_monthly.py <workdir> <year>   (needs openpyxl)
Dataset: https://www.tesourotransparente.gov.br/ckan/dataset/despesas-e-transferencias-totais
Writes <workdir>/cache/dtt_monthly_<year>.json.gz with PAGAMENTOS_TOTAIS summed by
[month, orgao, unidade orcamentaria, grupo, modalidade, acao, Categoria_RTN], plus the CKAN resource metadata
(last_modified is the date the Tesouro last published the file). Same reading as agg_dtt_full.py; the month is kept
so a period shorter than the year (e.g. 2026 Jan to Jul) can be cut. The XLSX is deleted at the end.
"""
import collections, gzip, json, os, subprocess, sys, zipfile
import openpyxl

W, Y = sys.argv[1], sys.argv[2]
API = "https://www.tesourotransparente.gov.br/ckan/api/3/action/package_show?id=despesas-e-transferencias-totais"
os.makedirs(os.path.join(W, "cache"), exist_ok=True); os.makedirs(os.path.join(W, "raw"), exist_ok=True)
res = json.loads(subprocess.run(["curl", "-sS", "-m", "60", "-A", "Mozilla/5.0", API], capture_output=True, check=True).stdout)["result"]["resources"]
meta = {r["name"][-4:]: r for r in res if r["name"].startswith("Despesas e Transfer")}
r = meta[Y]
out = os.path.join(W, "cache", f"dtt_monthly_{Y}.json.gz")
xp = os.path.join(W, "raw", f"dtt_{Y}.xlsx")
for attempt in range(5):
    if os.path.exists(xp) and zipfile.is_zipfile(xp): break
    subprocess.run(["curl", "-sS", "-L", "-m", "1500", "-A", "Mozilla/5.0", "-o", xp, r["url"]])
else:
    raise SystemExit("download failed")
wb = openpyxl.load_workbook(xp, read_only=True)
sheet = next(n for n in wb.sheetnames if next(wb[n].iter_rows(max_row=1, values_only=True))[0] == "ID_ANO")
rows = wb[sheet].iter_rows(values_only=True)
header = [str(c) for c in next(rows)]
ix = {c.upper(): i for i, c in enumerate(header)}
g = lambda c: ix[c.upper()]
im, io_, ion, iuo, iuon, ign, imo, iac, iacn, ipt, irtn = (g(c) for c in (
    "ID_MES", "ORGAO_CODIGO", "ORGAO_DESCRICAO", "UNIDADE_ORCAMENTARIA_CODIGO", "UNIDADE_ORCAMENTARIA_DESCRICAO",
    "ID_GRUPO_DESPESA_NADE", "CO_MOAP_NADE", "ACAO", "NO_ACAO", "PAGAMENTOS_TOTAIS", "Categoria_RTN"))
agg = collections.defaultdict(float)
names = {"orgao": {}, "uo": {}, "acao": {}}
n = 0
for row in rows:
    if row[im] is None: continue
    n += 1
    k = (int(row[im]), str(row[io_]), str(row[iuo]), str(row[ign]), str(row[imo]), str(row[iac]), str(row[irtn]))
    agg[k] += row[ipt] or 0.0
    names["orgao"].setdefault(k[1], row[ion]); names["uo"].setdefault(k[2], row[iuon]); names["acao"].setdefault(k[5], row[iacn])
wb.close(); size = os.path.getsize(xp); os.remove(xp)
with gzip.open(out, "wt", encoding="utf-8") as f:
    json.dump({"year": Y, "url": r["url"], "resource_last_modified": r.get("last_modified"), "resource_created": r.get("created"),
               "sheet": sheet, "header": header, "rows": n, "xlsx_bytes": size,
               "columns": ["month", "orgao", "uo", "grupo", "modalidade", "acao", "categoria_rtn", "pagamentos_totais"],
               "agg": [list(k) + [round(v, 2)] for k, v in agg.items() if v], "names": names}, f, ensure_ascii=False)
print(Y, "rows", n, "months", sorted({k[0] for k in agg}), "total_bi", round(sum(agg.values()) / 1e9, 2), "last_modified", r.get("last_modified"), flush=True)

"""Load the interest the União paid in cash (D121): Tesouro, base "Despesas e Transferências Totais da União",
grupo de natureza de despesa 2 (Juros e Encargos da Dívida), by month, one year file at a time.

Usage: .venv/bin/python pipelines/ledger/load_dtt_interest.py [YEAR ...]
Each file (23 to 38 MB) is read once, hashed while it streams, summed and deleted before the next one starts.
The money the Tesouro passes to the Banco Central (action 0669) and intra-budget payments stay inside the União:
internal_flow, never a flow to the creditors."""
import json
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal

import openpyxl

import db
import registry
import sources
from load_rtn import flow_id
from money import exact, to_cents

NEED = ("ID_ANO", "ID_MES", "ORGAO_CODIGO", "UNIDADE_ORCAMENTARIA_CODIGO", "ID_GRUPO_DESPESA_NADE", "CO_MOAP_NADE",
        "ACAO", "NO_ACAO", "DESPESAS_PAGAS", "RESTOS_A_PAGAR_PAGOS", "PAGAMENTOS_TOTAIS")
PARTS = {"paid": ("DESPESAS_PAGAS", "despesas do ano pagas"), "rap": ("RESTOS_A_PAGAR_PAGOS", "restos a pagar pagos")}
EPS = Decimal("0.001")


def resources():
    with registry.fetch(sources.DTT_CKAN, timeout=60) as fx:
        items = json.loads(fx.path.read_bytes())["result"]["resources"]
    return {int(x["name"][-4:]): x for x in items if x["name"].startswith("Despesas e Transferências da União - ")}


def read_interest(path, year):
    """{(key, period): {"paid": Decimal, "rap": Decimal}} and {key: action name}, group 2 only."""
    wb = openpyxl.load_workbook(path, read_only=True)
    sheet = next(n for n in wb.sheetnames if next(wb[n].iter_rows(max_row=1, values_only=True))[0] == "ID_ANO")
    rows = wb[sheet].iter_rows(values_only=True)
    header = [str(c).strip().upper() for c in next(rows)]
    missing = [c for c in NEED if c not in header]
    if missing:
        raise SystemExit(f"DTT {year}: columns missing {missing}")
    ix = {c: header.index(c) for c in NEED}
    agg, names = defaultdict(lambda: {"paid": Decimal(0), "rap": Decimal(0)}), {}
    for r in rows:
        if r[ix["ID_MES"]] is None or str(r[ix["ID_GRUPO_DESPESA_NADE"]]).strip() != sources.DTT_INTEREST_GROUP:
            continue
        if int(r[ix["ID_ANO"]]) != year:
            raise SystemExit(f"DTT {year}: a row of year {r[ix['ID_ANO']]}")
        p = date(year, int(r[ix["ID_MES"]]), 1)
        action = str(r[ix["ACAO"]]).strip().zfill(4)
        key = ".".join([str(r[ix["ORGAO_CODIGO"]]).strip(), str(r[ix["UNIDADE_ORCAMENTARIA_CODIGO"]]).strip(), action,
                        str(r[ix["CO_MOAP_NADE"]]).strip()])
        paid, rap, total = (exact(r[ix[c]] or 0) for c in ("DESPESAS_PAGAS", "RESTOS_A_PAGAR_PAGOS", "PAGAMENTOS_TOTAIS"))
        if abs(paid + rap - total) > EPS:
            raise SystemExit(f"DTT {year} {key} {p}: despesas pagas + restos a pagar pagos != pagamentos totais")
        agg[(key, p)]["paid"] += paid
        agg[(key, p)]["rap"] += rap
        names[key] = str(r[ix["NO_ACAO"]]).strip()
    wb.close()
    return agg, names


def build_rows(agg, names, source_id, dataset):
    published, flows, internal = [], [], []
    for (key, p), parts in sorted(agg.items()):
        if p > sources.LAST:
            continue
        action, modality = key.split(".")[2], key.split(".")[3]
        for part, (column, label) in PARTS.items():
            line_key, e = f"{key}|{part}", parts[part]
            cents = to_cents(e)
            published.append((source_id, line_key, p, "leaf", f"{names[key]} ({label})", e, cents, 1))
            if not cents:
                continue
            if action in sources.DTT_INTERNAL_ACTIONS:
                internal.append((source_id, line_key, None, p, "tesouro", sources.DTT_INTERNAL_ACTIONS[action], cents))
            elif modality == sources.DTT_INTRA_MODALITY:
                internal.append((source_id, line_key, None, p, "tesouro", "intra", cents))
            else:
                flows.append((flow_id(dataset, line_key, p, "gov:uniao", "fin:credores", "interest"), p, "gov:uniao",
                              "fin:credores", "interest", cents, "cash", source_id, line_key, "direct",
                              "no_state_endpoint_v1", None))
    return published, flows, internal


def run(con, years=sources.DTT_YEARS):
    res, out = resources(), []
    for year in years:
        r, dataset = res[year], f"tesouro.dtt.{year}"
        with registry.fetch(r["url"], suffix=".xlsx", timeout=1500) as fx:
            con.begin()
            try:
                source_id, new = registry.register(
                    con, dataset, fx, title=r["name"], publisher="Secretaria do Tesouro Nacional",
                    page_url=sources.DTT_PAGE, published_on=datetime.fromisoformat(r["last_modified"]).date(),
                    published_note="data de atualização do arquivo no portal (CKAN last_modified)",
                    loader="load_dtt_interest.py")
                if new:
                    db.insert(con, *build_rows(*read_interest(fx.path, year), source_id, dataset))
                con.commit()
            except BaseException:
                con.rollback()
                raise
        out.append({"dataset": dataset, "source_id": source_id, "new": new, "sha256": fx.sha256, "bytes": fx.bytes})
        print("dtt", out[-1], flush=True)
    return out


if __name__ == "__main__":
    years = tuple(int(a) for a in sys.argv[1:]) or sources.DTT_YEARS
    run(db.connect(), years)

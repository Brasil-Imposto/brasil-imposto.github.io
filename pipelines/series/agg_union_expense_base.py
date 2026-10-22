"""Aggregate the Tesouro Nacional "Despesas e Transferencias Totais da Uniao" yearly XLSX (criterio valor pago,
the base behind RTN table 1.5), keep what the federal payroll series needs, delete each raw file.

Usage: agg_union_expense_base.py <workdir> <year>[,<year>...]   (needs openpyxl)
Dataset: https://www.tesourotransparente.gov.br/ckan/dataset/despesas-e-transferencias-totais (resource URLs are read
from the CKAN API at run time). For every year it writes <workdir>/cache/dtt_YYYY.json with:
  gnd1:   rows of grupo de despesa 1, summed by [month, orgao, unidade orcamentaria, modalidade, acao, Categoria_RTN]
          -> PAGAMENTOS_TOTAIS (despesas pagas + restos a pagar pagos)
  fcdf:   rows of unidade orcamentaria 73901 (FCDF), any group, summed by [month, grupo, modalidade, acao, Categoria_RTN]
  rtn:    all rows summed by [month, Categoria_RTN] (to compare with RTN table 1.5)
  names:  code -> description for orgaos, unidades orcamentarias and acoes
"""
import collections, json, os, subprocess, sys, urllib.request, zipfile
import openpyxl

W = sys.argv[1]
YEARS = sys.argv[2].split(",")
API = "https://www.tesourotransparente.gov.br/ckan/api/3/action/package_show?id=despesas-e-transferencias-totais"
os.makedirs(os.path.join(W, "cache"), exist_ok=True); os.makedirs(os.path.join(W, "raw"), exist_ok=True)
res = json.load(urllib.request.urlopen(urllib.request.Request(API, headers={"User-Agent": "Mozilla/5.0"}), timeout=60))["result"]["resources"]
urls = {r["name"][-4:]: r["url"] for r in res if r["name"].startswith("Despesas e Transfer")}

for y in YEARS:
    out = os.path.join(W, "cache", f"dtt_{y}.json")
    if os.path.exists(out): continue
    xp = os.path.join(W, "raw", f"dtt_{y}.xlsx")
    for attempt in range(5):
        if os.path.exists(xp) and zipfile.is_zipfile(xp): break  # an XLSX is a ZIP: a cut download fails this test
        subprocess.run(["curl", "-sS", "-L", "-m", "900", "-A", "Mozilla/5.0", "-o", xp, urls[y]])
    else:
        raise SystemExit(f"{y}: download failed")
    wb = openpyxl.load_workbook(xp, read_only=True)
    # The data sheet is called "base" in some years and "<year>" in others: take the one whose header starts with ID_ANO.
    sheet = next(n for n in wb.sheetnames if next(wb[n].iter_rows(max_row=1, values_only=True))[0] == "ID_ANO")
    rows = wb[sheet].iter_rows(values_only=True)
    header = next(rows)
    ix = {str(c).upper(): i for i, c in enumerate(header)}  # column case changes between years (Categoria_RTN, CATEGORIA_RTN)
    def g(c):
        if c.upper() not in ix: raise SystemExit(f"{y}: column {c} missing; header = {header}")
        return ix[c.upper()]
    im, io_, ion, iuo, iuon, ign, imo, iac, iacn, ipt, irtn = (g(c) for c in (
        "ID_MES", "ORGAO_CODIGO", "ORGAO_DESCRICAO", "UNIDADE_ORCAMENTARIA_CODIGO", "UNIDADE_ORCAMENTARIA_DESCRICAO",
        "ID_GRUPO_DESPESA_NADE", "CO_MOAP_NADE", "ACAO", "NO_ACAO", "PAGAMENTOS_TOTAIS", "Categoria_RTN"))
    iyr = g("ID_ANO")
    gnd1, fcdf, rtn = (collections.defaultdict(float) for _ in range(3))
    names = {"orgao": {}, "uo": {}, "acao": {}}
    n = 0; years_seen = collections.Counter()
    for row in rows:
        if row[im] is None: continue
        n += 1; years_seen[row[iyr]] += 1
        v = row[ipt] or 0.0
        m = int(row[im]); cat = row[irtn] or ""
        rtn[(m, cat)] += v
        uo = str(row[iuo]); grp = str(row[ign])
        if uo == "73901":
            fcdf[(m, grp, str(row[imo]), str(row[iac]), cat)] += v
            names["acao"].setdefault(str(row[iac]), row[iacn])
        if grp == "1":
            gnd1[(m, str(row[io_]), uo, str(row[imo]), str(row[iac]), cat)] += v
            names["orgao"].setdefault(str(row[io_]), row[ion]); names["uo"].setdefault(uo, row[iuon])
            names["acao"].setdefault(str(row[iac]), row[iacn])
    wb.close(); size = os.path.getsize(xp); os.remove(xp)
    r = lambda d: [list(k) + [round(v, 2)] for k, v in d.items() if v]
    json.dump({"year": y, "url": urls[y], "sheet": sheet, "rows": n, "xlsx_bytes": size, "id_ano_values": {str(k): c for k, c in years_seen.items()},
               "gnd1": r(gnd1), "fcdf": r(fcdf), "rtn": r(rtn), "names": names}, open(out, "w"), ensure_ascii=False)
    print(y, "rows", n, "gnd1_bi", round(sum(gnd1.values()) / 1e9, 2), "fcdf_bi", round(sum(fcdf.values()) / 1e9, 2), flush=True)
print("DONE", flush=True)

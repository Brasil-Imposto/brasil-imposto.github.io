# Conferências e diagnósticos citados no relatório (não gera o JSON).
# Entradas: arrecadacao-estado.csv, serie-1994-2025.xlsx, tt/*.csv, transf_agg_2024.json, desp_agg_2024.json.
# Requer openpyxl (seção de setembro de 2024).
import sys, os, re, json, collections
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
from parse_rfb import rows, h, f, agg
from parse_tt import parse, annual

print('== 1. Formatos numéricos por ano no CSV da Receita')
fmt=collections.Counter()
for r in rows[1:]:
    for x in r[3:]:
        x=x.strip()
        if not x: continue
        if re.fullmatch(r'-?\d+',x): k='int'
        elif re.fullmatch(r'-?\d+,\d+',x): k='dec_comma'
        elif re.fullmatch(r'-?\d{1,3}(\.\d{3})+',x): k='thousand_dot'
        elif re.fullmatch(r'-?\d{1,3}(,\d{3})+',x): k='thousand_comma'
        else: k='other'
        fmt[(r[0],k)]+=1
for y in sorted(set(y for y,_ in fmt)): print(y,{k:v for (yy,k),v in fmt.items() if yy==y})

print('== 2. Participação de DF/SP/RJ/MG na arrecadação federal, por ano')
for y in ['2019','2020','2021','2022','2023','2024','2025']:
    a=agg(y); t=sum(sum(v.values()) for v in a.values())
    print(y,'BR %.1f bi'%(t/1e9),' '.join('%s %.1f%%'%(u,100*sum(a[u].values())/t) for u in ['DF','SP','RJ','MG']))

print('== 3. DF mês a mês (R$ bi) e DF/SP por tributo 2024 x 2025')
for y in ['2024','2025','2026']:
    print(y,[(r[1],round(sum(f(x) for x in r[3:])/1e9,1)) for r in rows[1:] if r[0]==y and r[2]=='DF'])
a4=agg('2024'); a5=agg('2025')
for c in h[3:]:
    b4=sum(v[c] for v in a4.values()); b5=sum(v[c] for v in a5.values())
    if b4 or b5: print('%-60s DF %7.1f -> %7.1f | SP %7.1f -> %7.1f'%(c[:60],a4['DF'][c]/1e9,a5['DF'][c]/1e9,a4['SP'][c]/1e9,a5['SP'][c]/1e9))

print('== 4. Setembro de 2024: Imposto de Importação e de Exportação, UF x nacional (R$ mi)')
import openpyxl
ws=openpyxl.load_workbook('serie-1994-2025.xlsx',data_only=True)['2024']
ii=[row for row in ws.iter_rows(values_only=True) if row[0]=='IMPOSTO SOBRE IMPORTAÇÃO'][0]
ie=[row for row in ws.iter_rows(values_only=True) if row[0]=='IMPOSTO SOBRE EXPORTAÇÃO'][0]
meses=['Janeiro','Fevereiro','Março','Abril','Maio','Junho','Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']
m=collections.defaultdict(float); mi=collections.defaultdict(float)
for r in rows[1:]:
    if r[0]=='2024': m[r[1]]+=f(r[4]); mi[r[1]]+=f(r[3])
for i,mm in enumerate(meses):
    print('%-10s UF_II %8.1f NAT_II %8.1f | UF_IE %8.1f NAT_IE %8.3f'%(mm,mi[mm]/1e6,ii[i+1],m[mm]/1e6,ie[i+1]))

print('== 5. Portal da Transparência x Tesouro (2024)')
tr=json.load(open('transf_agg_2024.json'))
for fn,a in [('fpe','0044'),('fpm','0045'),('ipi-exp','0046'),('itr','006M')]:
    d,_,_=parse(f'tt/{fn}.csv'); t=sum(annual(d,2024).values()); p=sum(r[-1] for r in tr if r[5]==a)
    print(fn,'tesouro %.2f portal %.2f diff %.2f'%(t,p,p-t))
de,_,_=parse('tt/lc176est.csv'); dm,_,_=parse('tt/lc176mun.csv')
print('lc176 tesouro %.2f portal (sem UF) %.2f'%(sum(annual(de,2024).values())+sum(annual(dm,2024).values()),sum(r[-1] for r in tr if r[5]=='00SE')))

print('== 6. Execução da despesa 2024: quanto do gasto direto tem UF')
d=json.load(open('desp_agg_2024.json'))
V=lambda r:r[9]+r[10]
direct=lambda r:r[4] in ('90','91') and r[3] in ('1','3','4','5')
tot=sum(V(r) for r in d if direct(r)); wuf=sum(V(r) for r in d if direct(r) and r[0])
print('gasto direto primário %.2f bi; com UF %.2f bi (%.1f%%)'%(tot/1e9,wuf/1e9,100*wuf/tot))
g1=sum(V(r) for r in d if r[3]=='1'); g1u=sum(V(r) for r in d if r[3]=='1' and r[0])
print('pessoal (grupo 1) %.2f bi; com UF %.2f bi (%.1f%%)'%(g1/1e9,g1u/1e9,100*g1u/g1))
print('pessoal por UF:',sorted(((k or 'sem UF',round(sum(V(r) for r in d if r[3]=='1' and r[0]==k)/1e9,2)) for k in set(r[0] for r in d)),key=lambda t:-t[1])[:6])
loc=collections.Counter()
for r in d:
    if direct(r) and not r[0]: loc[r[8]]+=V(r)
print('localizadores sem UF (0001 = NACIONAL):',[(k,round(v/1e9,1)) for k,v in loc.most_common(6)])
for a in ['00SJ','00IN','00H5','8442','00H4','0581','20TP','0181','2867','214H','0179']:
    s=sum(V(r) for r in d if r[6]==a); su=sum(V(r) for r in d if r[6]==a and r[0])
    print(a,[r[7] for r in d if r[6]==a][0][:60],'%.2f bi, com UF %.2f'%(s/1e9,su/1e9))
